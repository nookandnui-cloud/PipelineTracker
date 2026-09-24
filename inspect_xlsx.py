#!/usr/bin/env python3
"""Inspect the pipeline workbook: exact headers, cell formats, sample rows."""
import sys
from openpyxl import load_workbook

XLSX = sys.argv[1] if len(sys.argv) > 1 else "Pipeline system team.xlsx"

wb = load_workbook(XLSX, data_only=True)
print("SHEETS:", wb.sheetnames)

ws = wb[wb.sheetnames[-1]] if len(wb.sheetnames) > 1 else wb.active
# find the Pipeline sheet
for name in wb.sheetnames:
    if "pipeline" in name.lower():
        ws = wb[name]
        break
print("USING SHEET:", ws.title, "dims:", ws.dimensions, "rows:", ws.max_row, "cols:", ws.max_column)

hdr = None
for r in range(1, 25):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    if any(isinstance(v, str) and v.strip().lower() in ("project code", "project name") or
           (isinstance(v, str) and "project code" in v.lower()) for v in vals):
        hdr = r
        break
print("HEADER ROW:", hdr)
for c in range(1, ws.max_column + 1):
    cell = ws.cell(hdr, c)
    print(f"  col {cell.column_letter} ({c}): {cell.value!r}")

print("\n--- SAMPLE DATA ROWS (cell by cell) ---")
for r in list(range(hdr + 1, hdr + 4)) + [hdr + 10, hdr + 40]:
    print(f"=== row {r}")
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(r, c)
        if cell.value is not None:
            print(f"  {cell.column_letter}{r}: {cell.value!r}   [fmt={cell.number_format!r}]")

# Collect distinct values for key columns to catch typos
print("\n--- DISTINCT VALUES ---")
def distinct(cidx, limit=40):
    seen = []
    for r in range(hdr + 1, ws.max_row + 1):
        v = ws.cell(r, cidx).value
        if v is not None and str(v).strip() != "" and v not in seen:
            seen.append(v)
    return seen[:limit]

for label, key in [("TEAM", "team"), ("STATUS(W/L/D)", "win/lost"), ("START", "start"), ("TARGET", "target")]:
    for c in range(1, ws.max_column + 1):
        h = ws.cell(hdr, c).value
        if isinstance(h, str) and h.strip().lower().startswith(key):
            print(f"{label}: {distinct(c)}")
            break

# numeric column between Target and Project name: check formats
print("\n--- NUMERIC 'Start Date?' COLUMN FORMATS ---")
for c in range(1, ws.max_column + 1):
    h = ws.cell(hdr, c).value
    if isinstance(h, str) and "start date" in h.strip().lower():
        col = c
        print("column:", ws.cell(hdr, c).column_letter, repr(h))
        n_int, n_dt, n_other, fmts = 0, 0, 0, set()
        samples = []
        for r in range(hdr + 1, ws.max_row + 1):
            v = ws.cell(r, c).value
            if v is None:
                continue
            fmts.add(ws.cell(r, c).number_format)
            if isinstance(v, (int, float)):
                n_int += 1
                if len(samples) < 8:
                    samples.append((r, v, ws.cell(r, c).number_format))
            elif hasattr(v, "year"):
                n_dt += 1
            else:
                n_other += 1
                if len(samples) < 12:
                    samples.append((r, repr(v), ws.cell(r, c).number_format))
        print(f"ints={n_int} dates={n_dt} other={n_other}")
        print("formats:", fmts)
        print("samples:", samples)
        break
